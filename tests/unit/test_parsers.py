"""
Unit Tests for Multi-Format Parsers (PDF, DOCX, TXT/MD)
Covers multi-page PDFs, headings, tables, merged cells, BOM stripping, multilingual text, and corrupt files.
"""

import pytest
import tempfile
import pathlib
import docx
import fitz
from app.services.ingestion.parsers import (
    PyMuPDFParser,
    DocxParser,
    TxtParser,
    ParserRegistry,
    default_parser_registry,
)
from app.services.ingestion.exceptions import (
    UnsupportedFileTypeError,
    FileNotFoundIngestionError,
    CorruptedFileError,
)

@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as tmp:
        yield pathlib.Path(tmp)

# --- TXT / Markdown Parser Tests ---

def test_txt_parser_standard_and_bom(temp_dir: pathlib.Path):
    parser = TxtParser()

    # 1. Standard Markdown
    md_file = temp_dir / "sample.md"
    md_content = "# Academic Regulations\n\n## 1.0 Minimum Attendance\n\nStudents must maintain 75% attendance.\n\n## 2.0 Fee Payment\n\nAll fees must be cleared."
    md_file.write_text(md_content, encoding="utf-8")

    parsed = parser.parse(md_file, doc_id="DOC-TXT-001", category="academic_regulations")
    assert parsed.doc_id == "DOC-TXT-001"
    assert parsed.file_type == "md"
    assert parsed.page_count == 1
    assert "Academic Regulations" in parsed.normalized_text
    assert len(parsed.sections) == 3
    assert parsed.sections[0]["title"] == "Academic Regulations"
    assert parsed.sections[1]["title"] == "1.0 Minimum Attendance"
    assert parsed.detected_language == "en"

    # 2. UTF-8 with BOM
    bom_file = temp_dir / "bom_sample.txt"
    bom_bytes = b"\xef\xbb\xbfSection 1\nHostel Curfew is 9:30 PM."
    bom_file.write_bytes(bom_bytes)

    parsed_bom = parser.parse(bom_file, doc_id="DOC-TXT-002", category="hostel")
    assert "Hostel Curfew is 9:30 PM." in parsed_bom.normalized_text
    assert any(w.warning_code == "BOM_STRIPPED" for w in parsed_bom.warnings)

def test_txt_parser_multilingual_and_windows_newlines(temp_dir: pathlib.Path):
    parser = TxtParser()
    crlf_file = temp_dir / "crlf_multilingual.txt"
    content = (
        "1.0 Hindi Attendance Rules\r\n"
        "छात्रों को 75% उपस्थिति बनाए रखनी होगी।\r\n\r\n"
        "2.0 Kannada Rules\r\n"
        "ಹಾಜರಾತಿ ನಿಯಮಗಳು ಕಡ್ಡಾಯವಾಗಿವೆ.\r\n"
    )
    crlf_file.write_bytes(content.encode("utf-8"))

    parsed = parser.parse(crlf_file, doc_id="DOC-TXT-003", category="attendance")
    assert "\r" not in parsed.normalized_text
    assert "छात्रों को 75% उपस्थिति बनाए रखनी होगी।" in parsed.normalized_text
    assert "ಹಾಜರಾತಿ ನಿಯಮಗಳು ಕಡ್ಡಾಯವಾಗಿವೆ." in parsed.normalized_text
    assert parsed.detected_script in ("devanagari", "kannada", "mixed")

def test_txt_parser_empty_file(temp_dir: pathlib.Path):
    parser = TxtParser()
    empty_file = temp_dir / "empty.txt"
    empty_file.write_bytes(b"")

    parsed = parser.parse(empty_file, doc_id="DOC-TXT-004", category="general")
    assert parsed.raw_text == ""
    assert any(w.warning_code == "EMPTY_FILE" for w in parsed.warnings)

# --- DOCX Parser Tests ---

def test_docx_parser_with_styles_and_table(temp_dir: pathlib.Path):
    parser = DocxParser()
    docx_file = temp_dir / "test_policy.docx"

    doc = docx.Document()
    doc.add_heading("Autonomous Examination Guidelines", level=1)
    doc.add_paragraph("All students must carry their hall ticket.")
    doc.add_heading("Section 2: Grading Scale", level=2)
    doc.add_paragraph("Grades are assigned on an absolute scale.")

    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Grade"
    table.rows[0].cells[1].text = "Marks Range"
    table.rows[1].cells[0].text = "S"
    table.rows[1].cells[1].text = "90-100"
    table.rows[2].cells[0].text = "A"
    table.rows[2].cells[1].text = "80-89"

    doc.save(docx_file)

    parsed = parser.parse(docx_file, doc_id="DOC-DOCX-001", category="examination_guidelines")
    assert parsed.doc_id == "DOC-DOCX-001"
    assert parsed.file_type == "docx"
    assert parsed.page_count == 1
    assert "Autonomous Examination Guidelines" in parsed.normalized_text
    assert "Section 2: Grading Scale" in parsed.normalized_text
    assert "[TABLE]" in parsed.normalized_text
    assert "90-100" in parsed.normalized_text
    assert len(parsed.sections) == 2
    assert parsed.detected_language == "en"

def test_docx_parser_merged_cells_and_multilingual(temp_dir: pathlib.Path):
    parser = DocxParser()
    docx_file = temp_dir / "merged_table.docx"

    doc = docx.Document()
    doc.add_heading("ಕಾಲೇಜು ಶುಲ್ಕ ರಚನೆ (College Fee Structure)", level=1)
    doc.add_paragraph("ವಿದ್ಯಾರ್ಥಿ ವೇತನ ಮತ್ತು ಶುಲ್ಕ ವಿವರಗಳು.")

    # Create table with merged cells
    table = doc.add_table(rows=2, cols=3)
    # Merge row 0, cell 0 and cell 1
    cell_a = table.rows[0].cells[0]
    cell_b = table.rows[0].cells[1]
    cell_a.merge(cell_b)
    cell_a.text = "ವರ್ಗ (Category Details)"
    table.rows[0].cells[2].text = "ಶುಲ್ಕ (Fee)"

    table.rows[1].cells[0].text = "ಜನರಲ್ (General)"
    table.rows[1].cells[1].text = "ಮೆರಿಟ್ (Merit)"
    table.rows[1].cells[2].text = "₹50,000"

    doc.save(docx_file)

    parsed = parser.parse(docx_file, doc_id="DOC-DOCX-002", category="scholarships")
    assert "ಕಾಲೇಜು ಶುಲ್ಕ ರಚನೆ" in parsed.normalized_text
    assert "[TABLE]" in parsed.normalized_text
    assert "₹50,000" in parsed.normalized_text
    assert parsed.detected_language == "kn" or parsed.detected_script in ("kannada", "mixed")

def test_docx_parser_empty_and_corrupt(temp_dir: pathlib.Path):
    parser = DocxParser()
    
    # 1. Empty document
    empty_doc = docx.Document()
    empty_path = temp_dir / "empty.docx"
    empty_doc.save(empty_path)

    parsed = parser.parse(empty_path, doc_id="DOC-DOCX-003", category="general")
    assert any(w.warning_code == "EMPTY_DOCUMENT" for w in parsed.warnings)

    # 2. Corrupted file bytes
    corrupt_path = temp_dir / "corrupted.docx"
    corrupt_path.write_bytes(b"PK\x03\x04not_a_real_docx")
    with pytest.raises(CorruptedFileError):
        parser.parse(corrupt_path, doc_id="DOC-DOCX-004", category="general")

# --- PDF Parser Tests ---

def test_pdf_parser_multipage_and_empty_page_warning(temp_dir: pathlib.Path):
    parser = PyMuPDFParser()
    pdf_file = temp_dir / "multipage.pdf"

    doc = fitz.open()

    # Page 1: Heading & Text
    p1 = doc.new_page()
    p1.insert_text((50, 50), "SECTION 1.0 SCHOLARSHIPS", fontsize=14)
    p1.insert_text((50, 80), "Institutional Merit-cum-Means scheme provides 50% tuition fee waiver for eligible students.")

    # Page 2: Empty Page (triggers OCR warning)
    p2 = doc.new_page()

    # Page 3: Additional Content
    p3 = doc.new_page()
    p3.insert_text((50, 50), "SECTION 2.0 APPLICATION DEADLINES", fontsize=14)
    p3.insert_text((50, 80), "Applications must be submitted through the scholarship portal by October 15th.")

    doc.save(pdf_file)
    doc.close()

    parsed = parser.parse(pdf_file, doc_id="DOC-PDF-001", category="scholarships")
    assert parsed.doc_id == "DOC-PDF-001"
    assert parsed.file_type == "pdf"
    assert parsed.page_count == 3
    assert len(parsed.pages) == 3
    assert parsed.pages[0].page_number == 1
    assert parsed.pages[1].page_number == 2
    assert parsed.pages[2].page_number == 3

    assert any(w.warning_code == "OCR_REQUIRED_OR_TEXT_NOT_EXTRACTABLE" and w.page_number == 2 for w in parsed.warnings)
    assert "SECTION 1.0 SCHOLARSHIPS" in parsed.normalized_text
    assert "SECTION 2.0 APPLICATION DEADLINES" in parsed.normalized_text

def test_pdf_parser_corrupt_file(temp_dir: pathlib.Path):
    parser = PyMuPDFParser()
    corrupt_path = temp_dir / "corrupted.pdf"
    corrupt_path.write_bytes(b"%PDF-1.4\ncorrupted_binary_data_without_trailer")

    with pytest.raises(CorruptedFileError):
        parser.parse(corrupt_path, doc_id="DOC-PDF-002", category="general")

# --- Registry Tests ---

def test_parser_registry_resolution_and_errors(temp_dir: pathlib.Path):
    registry = ParserRegistry()

    assert isinstance(registry.get_parser_for_file("file.pdf"), PyMuPDFParser)
    assert isinstance(registry.get_parser_for_file("file.docx"), DocxParser)
    assert isinstance(registry.get_parser_for_file("file.txt"), TxtParser)
    assert isinstance(registry.get_parser_for_file("file.md"), TxtParser)

    with pytest.raises(UnsupportedFileTypeError):
        registry.get_parser_for_file("file.unsupported_xyz")

    with pytest.raises(UnsupportedFileTypeError):
        registry.get_parser_for_file("file_without_extension")
